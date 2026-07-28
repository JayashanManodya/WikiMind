"""Document parser supporting PDF, DOCX, TXT, and MD files."""

import io
import re
from typing import Dict, Any, List
import fitz  # PyMuPDF
import docx


def parse_document_bytes(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Parse document bytes into structured JSON representation containing pages and full_text.

    Args:
        file_bytes: Raw byte contents of the uploaded file.
        filename: Original filename to infer document type (.pdf, .docx, .txt, .md).

    Returns:
        Dict with keys: filename, total_pages, pages (list of {page_number, text}), full_text.
    """
    ext = filename.lower().split(".")[-1] if "." in filename else ""

    if ext == "pdf":
        return _parse_pdf(file_bytes, filename)
    elif ext in ["docx", "doc"]:
        return _parse_docx(file_bytes, filename)
    elif ext in ["txt", "md"]:
        return _parse_text(file_bytes, filename)
    else:
        try:
            return _parse_text(file_bytes, filename)
        except Exception as e:
            raise ValueError(f"Unsupported file type .{ext}: {str(e)}")


def _parse_pdf(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Parse PDF file bytes using PyMuPDF (fitz)."""
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    pages: List[Dict[str, Any]] = []
    full_text_parts: List[str] = []

    for page_num in range(len(doc)):
        page = doc[page_num]
        text = page.get_text("text") or ""
        pages.append({
            "page_number": page_num + 1,
            "text": text
        })
        full_text_parts.append(text)

    doc.close()
    full_text = "\n\n".join(full_text_parts).strip()

    return {
        "filename": filename,
        "total_pages": len(pages),
        "pages": pages,
        "full_text": full_text
    }


def _parse_docx(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Parse DOCX file bytes using python-docx."""
    buffer = io.BytesIO(file_bytes)
    doc = docx.Document(buffer)
    
    text_parts: List[str] = []

    for p in doc.paragraphs:
        if p.text.strip():
            text_parts.append(p.text.strip())

    for table in doc.tables:
        table_lines: List[str] = []
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells]
            table_lines.append(" | ".join(row_cells))
        if table_lines:
            text_parts.append("\n".join(table_lines))

    full_text = "\n\n".join(text_parts).strip()

    return {
        "filename": filename,
        "total_pages": 1,
        "pages": [{"page_number": 1, "text": full_text}],
        "full_text": full_text
    }


def _parse_text(file_bytes: bytes, filename: str) -> Dict[str, Any]:
    """Parse UTF-8 plain text or Markdown file bytes."""
    try:
        content = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        content = file_bytes.decode("latin-1", errors="replace")

    return {
        "filename": filename,
        "total_pages": 1,
        "pages": [{"page_number": 1, "text": content}],
        "full_text": content
    }
