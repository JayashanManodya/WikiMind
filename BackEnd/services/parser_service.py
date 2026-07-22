import os
import json
import fitz  # PyMuPDF
import docx
from pathlib import Path
from datetime import datetime, timezone
from typing import List, Optional
from fastapi import HTTPException, status

from BackEnd.schemas import ParsedPage, ParsedDocumentResponse
from BackEnd.services.metadata_service import metadata_service

class ParserService:
    def __init__(self, parsed_dir: str = "./storage/parsed"):
        self.parsed_dir = Path(parsed_dir)
        self.parsed_dir.mkdir(parents=True, exist_ok=True)

    def parse_pdf(self, file_path: Path) -> List[ParsedPage]:
        """Extract page-by-page text from PDF using PyMuPDF (fitz), preserving page order and layout"""
        pages: List[ParsedPage] = []
        try:
            doc = fitz.open(str(file_path))
            for page_num in range(len(doc)):
                page = doc.load_page(page_num)
                text = page.get_text("text") or ""
                text = text.strip()
                
                pages.append(ParsedPage(
                    page_number=page_num + 1,
                    text=text,
                    char_count=len(text)
                ))
            doc.close()
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"PyMuPDF PDF parsing error: {str(e)}"
            )
        return pages

    def parse_docx(self, file_path: Path) -> List[ParsedPage]:
        """Extract sequential text paragraphs and table contents from DOCX using python-docx"""
        pages: List[ParsedPage] = []
        text_elements: List[str] = []
        try:
            document = docx.Document(str(file_path))
            
            # Extract paragraphs
            for paragraph in document.paragraphs:
                p_text = paragraph.text.strip()
                if p_text:
                    text_elements.append(p_text)
                    
            # Extract tables
            for table in document.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        text_elements.append(row_text)

            full_docx_text = "\n\n".join(text_elements)
            pages.append(ParsedPage(
                page_number=1,
                text=full_docx_text,
                char_count=len(full_docx_text)
            ))
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"DOCX parsing error: {str(e)}"
            )
        return pages

    def parse_txt(self, file_path: Path) -> List[ParsedPage]:
        """Extract plain text or markdown file content with UTF-8 encoding handling"""
        try:
            with open(file_path, "r", encoding="utf-8", errors="replace") as f:
                content = f.read().strip()
                
            pages = [ParsedPage(
                page_number=1,
                text=content,
                char_count=len(content)
            )]
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"TXT/MD parsing error: {str(e)}"
            )
        return pages

    def parse_document(self, file_id: str) -> ParsedDocumentResponse:
        """
        Main orchestrator: Fetches document metadata by ID, selects appropriate parser by extension,
        saves parsed JSON artifact to storage/parsed/<file_id>_parsed.json, and updates status.
        """
        metadata = metadata_service.get_metadata_by_id(file_id)
        if not metadata:
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Document with file_id '{file_id}' not found."
            )

        file_path = Path(metadata.file_path)
        if not file_path.exists():
            raise HTTPException(
                status_code=status.HTTP_404_NOT_FOUND,
                detail=f"Source document file not found at path '{file_path}'."
            )

        ext = file_path.suffix.lower()
        if ext == ".pdf":
            parsed_pages = self.parse_pdf(file_path)
        elif ext == ".docx":
            parsed_pages = self.parse_docx(file_path)
        elif ext in {".txt", ".md"}:
            parsed_pages = self.parse_txt(file_path)
        else:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Unsupported file format '{ext}' for parsing."
            )

        # Assemble full text with clear page dividers
        page_texts = [f"--- Page {p.page_number} ---\n{p.text}" for p in parsed_pages]
        full_text = "\n\n".join(page_texts)

        parsed_response = ParsedDocumentResponse(
            file_id=metadata.file_id,
            original_filename=metadata.original_filename,
            total_pages=len(parsed_pages),
            full_text=full_text,
            pages=parsed_pages,
            parsed_at=datetime.now(timezone.utc).isoformat(),
            status="parsed"
        )

        # Persist JSON artifact
        json_path = self.parsed_dir / f"{file_id}_parsed.json"
        with open(json_path, "w", encoding="utf-8") as f:
            f.write(parsed_response.model_dump_json(indent=2))

        # Update metadata status
        metadata.status = "parsed"
        metadata_service.save_metadata(metadata)

        return parsed_response

    def get_parsed_document(self, file_id: str) -> Optional[ParsedDocumentResponse]:
        """Retrieve stored parsed document response JSON if available"""
        json_path = self.parsed_dir / f"{file_id}_parsed.json"
        if not json_path.exists():
            return None
        try:
            with open(json_path, "r", encoding="utf-8") as f:
                data = json.load(f)
                return ParsedDocumentResponse(**data)
        except Exception:
            return None

# Default singleton instance
parser_service = ParserService()
